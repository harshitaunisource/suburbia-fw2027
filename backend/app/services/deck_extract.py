"""
Pulls every embedded image out of an uploaded .pptx or .pdf deck --
these are trend/vendor catalogue files that are mostly a grid of product
photos with little or no text (per the actual use case: a buyer trend
deck or vendor catalogue is images "here and there", not a normal
slide/document with structured content).

Uses each image's actual embedded bytes (not a screenshot/render of the
slide/page), so quality matches whatever was originally placed in the
file, and small/logo-sized images can be filtered out before they ever
reach the (paid) AI classification step.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

from PIL import Image

# Below this, an embedded image is almost certainly a bullet icon, logo,
# or decorative divider, not a product photo -- skip it before it ever
# reaches AI classification (saves real money on OpenAI vision calls).
MIN_DIMENSION_PX = 120


class DeckExtractError(Exception):
    pass


@dataclass
class ExtractedImage:
    image_bytes: bytes
    content_type: str
    slide_number: int  # 1-indexed slide/page the image came from


def _passes_size_filter(image_bytes: bytes) -> bool:
    try:
        with Image.open(io.BytesIO(image_bytes)) as img:
            w, h = img.size
            return w >= MIN_DIMENSION_PX and h >= MIN_DIMENSION_PX
    except Exception:
        # If PIL can't even open it, it's not a usable product photo.
        return False


_EXT_TO_CONTENT_TYPE = {
    "png": "image/png",
    "jpeg": "image/jpeg",
    "jpg": "image/jpeg",
    "gif": "image/gif",
    "bmp": "image/bmp",
    "webp": "image/webp",
}


def extract_images_from_pptx(file_bytes: bytes) -> list[ExtractedImage]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        prs = Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DeckExtractError(f"Could not open this file as a .pptx: {exc}") from exc

    images: list[ExtractedImage] = []

    def _walk(shapes, slide_number):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                _walk(shape.shapes, slide_number)
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(shape, "image", None):
                try:
                    image = shape.image
                except Exception:
                    continue
                if not _passes_size_filter(image.blob):
                    continue
                ext = (image.ext or "png").lower()
                images.append(ExtractedImage(
                    image_bytes=image.blob,
                    content_type=_EXT_TO_CONTENT_TYPE.get(ext, "image/png"),
                    slide_number=slide_number,
                ))

    for i, slide in enumerate(prs.slides, start=1):
        _walk(slide.shapes, i)

    return images


def extract_images_from_pdf(file_bytes: bytes) -> list[ExtractedImage]:
    import pymupdf

    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise DeckExtractError(f"Could not open this file as a .pdf: {exc}") from exc

    images: list[ExtractedImage] = []
    try:
        for page_number in range(len(doc)):
            page = doc[page_number]
            for img in page.get_images(full=True):
                xref = img[0]
                try:
                    base_image = doc.extract_image(xref)
                except Exception:
                    continue
                if not _passes_size_filter(base_image["image"]):
                    continue
                ext = base_image.get("ext", "png").lower()
                images.append(ExtractedImage(
                    image_bytes=base_image["image"],
                    content_type=_EXT_TO_CONTENT_TYPE.get(ext, "image/png"),
                    slide_number=page_number + 1,
                ))
    finally:
        doc.close()

    return images


def extract_images(filename: str, file_bytes: bytes) -> list[ExtractedImage]:
    """Dispatches on file extension. Raises DeckExtractError for anything
    that isn't a .pptx or .pdf -- this feature is specifically for
    image-heavy decks, not a general file-upload endpoint."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pptx":
        return extract_images_from_pptx(file_bytes)
    if suffix == ".pdf":
        return extract_images_from_pdf(file_bytes)
    raise DeckExtractError(
        f"Unsupported file type '{suffix}' -- upload a .pptx or .pdf trend/catalogue deck."
    )